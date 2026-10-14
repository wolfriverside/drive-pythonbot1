import os
import re
import json
import csv
from urllib.parse import urlparse, unquote
from pypdf import PdfReader
from docx import Document
from core.security import is_safe_local_path, mask_error


def _read_text_with_encodings(filepath: str) -> str:
    """嘗試多種繁體中文與常見編碼讀取純文字檔案。"""
    encodings = ["utf-8-sig", "utf-8", "cp950", "big5", "gb18030", "latin1"]
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
        except Exception as e:
            raise e
    raise UnicodeDecodeError("multiple", b"", 0, 1, "無法以常見編碼 (UTF-8, CP950/Big5, GB18030) 解碼檔案")


def _read_csv_file(filepath: str) -> str:
    """以自適應編碼讀取 CSV 檔案，避免 UnicodeDecodeError。"""
    raw = _read_text_with_encodings(filepath)
    reader = csv.reader(raw.splitlines())
    lines = []
    for row in reader:
        if row:
            lines.append(", ".join(row))
    return "\n".join(lines)


def _read_excel_file(filepath: str, ext: str) -> str:
    """
    讀取 Excel 檔案 (.xlsx, .xls)。
    支援標準 XLSX、標準二進位 XLS (BIFF8)、以及台灣常見的「HTML/XML 表格偽裝成 .xls」。
    """
    ext = ext.lower()

    # 1. 處理 .xlsx 格式
    if ext == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
            output = []
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                rows_text = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
                    if row_vals:
                        rows_text.append(" | ".join(row_vals))
                if rows_text:
                    output.append(f"[工作表: {sheetname}]\n" + "\n".join(rows_text))
            wb.close()
            if output:
                return "\n\n".join(output)
            return "（Excel 試算表內容為空）"
        except Exception as e_xlsx:
            # 降級嘗試純文字或檢查是否副檔名命名錯誤
            pass

    # 2. 處理 .xls 格式（包含偽裝成 .xls 的 HTML / XML）
    # 檢查檔案頭部判斷真實格式
    try:
        with open(filepath, "rb") as f:
            header_bytes = f.read(2048)
    except Exception as e:
        return f"Error: 無法讀取檔案標頭 - {e}"

    header_lower = header_bytes.lower()

    # 情況 2A: 偽裝成 XLS 的 HTML 表格（政府/學校系統極其常見）
    if b"<html" in header_lower or b"<!doctype" in header_lower or b"<table" in header_lower:
        try:
            import warnings
            from bs4 import BeautifulSoup, XMLParsedAsHTMLWarning
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", category=XMLParsedAsHTMLWarning)
                html_text = _read_text_with_encodings(filepath)
                soup = BeautifulSoup(html_text, "html.parser")
                tables = soup.find_all("table")
                output_tables = []
                for idx, table in enumerate(tables):
                    rows_out = []
                    for tr in table.find_all("tr"):
                        cells = [td.get_text(strip=True) for td in tr.find_all(["td", "th"])]
                        if any(cells):
                            rows_out.append(" | ".join(cells))
                    if rows_out:
                        output_tables.append(f"[表格 {idx+1}]\n" + "\n".join(rows_out))
                if output_tables:
                    return "\n\n".join(output_tables)
                # 若無明確 table 標籤，萃取頁面文字
                from tools.web_scraper import extract_clean_text_and_tables
                clean_txt, _ = extract_clean_text_and_tables(html_text)
                if clean_txt:
                    return clean_txt
        except Exception:
            pass

    # 情況 2B: XML Spreadsheet 2003 格式 (開頭為 <?xml ... <Workbook)
    if b"<?xml" in header_lower:
        try:
            import warnings
            from bs4 import BeautifulSoup, XMLParsedAsHTMLWarning
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", category=XMLParsedAsHTMLWarning)
                xml_text = _read_text_with_encodings(filepath)
                soup = BeautifulSoup(xml_text, "html.parser")
                worksheets = soup.find_all(re.compile(r"worksheet", re.I))
                output_sheets = []
                for s in worksheets:
                    sname = s.get("ss:name") or s.get("name") or "Sheet"
                    rows_out = []
                    for row in s.find_all(re.compile(r"row", re.I)):
                        cells = []
                        for cell in row.find_all(re.compile(r"^cell$", re.I)):
                            txt = cell.get_text(strip=True)
                            if txt:
                                cells.append(txt)
                        if cells:
                            rows_out.append(" | ".join(cells))
                    if rows_out:
                        output_sheets.append(f"[工作表: {sname}]\n" + "\n".join(rows_out))
                if output_sheets:
                    return "\n\n".join(output_sheets)
                # 若無 worksheet 標籤，純文字萃取
                raw_txt = soup.get_text(separator="\n", strip=True)
                if raw_txt:
                    return raw_txt
        except Exception:
            pass
        return "（XML 試算表內容為空或無法解析）"

    # 情況 2C: 標準二進位 .xls (BIFF8) 格式，使用 xlrd 讀取
    try:
        import xlrd
        wb = xlrd.open_workbook(filepath, formatting_info=False)
        output = []
        for sname in wb.sheet_names():
            sheet = wb.sheet_by_name(sname)
            rows_text = []
            for r in range(sheet.nrows):
                row_vals = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    val = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            dt = xlrd.xldate_as_datetime(val, wb.datemode)
                            val = dt.strftime("%Y-%m-%d %H:%M:%S" if dt.hour or dt.minute else "%Y-%m-%d")
                        except Exception:
                            val = str(val)
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        if isinstance(val, float) and val.is_integer():
                            val = str(int(val))
                        else:
                            val = str(val)
                    elif val is None:
                        val = ""
                    else:
                        val = str(val).strip()
                    row_vals.append(val)
                # 忽略全空列
                if any(v for v in row_vals):
                    rows_text.append(" | ".join(row_vals))
            if rows_text:
                output.append(f"[工作表: {sname}]\n" + "\n".join(rows_text))
        if output:
            return "\n\n".join(output)
        return "（Excel 試算表內容為空）"
    except Exception as e_xlrd:
        # 情況 2D: 若 xlrd 失敗，嘗試當作 Tab/Comma 分隔的純文字試算表
        try:
            tsv_txt = _read_text_with_encodings(filepath)
            lines = [line.strip() for line in tsv_txt.splitlines() if line.strip()]
            if lines and ("\t" in lines[0] or "," in lines[0]):
                return "\n".join(lines)
        except Exception:
            pass
        return f"Error: 讀取 XLS 試算表時發生錯誤（{e_xlrd}）"


def read_local_file(filepath: str) -> str:
    """
    Reads a local file or downloads an online file (PDF, TXT, CSV, DOCX, XLSX, XLS, JSON) and extracts its content.
    Supports academic papers, journal articles, government documents, and spreadsheets from URLs.
    """
    temp_file_path = None
    if filepath.startswith("http://") or filepath.startswith("https://"):
        import requests
        import tempfile

        parsed_origin = urlparse(filepath)
        origin_url = f"{parsed_origin.scheme}://{parsed_origin.netloc}"

        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36",
            "Accept": "text/html,application/xhtml+xml,application/xml,application/pdf,application/vnd.ms-excel,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*;q=0.8",
            "Accept-Language": "zh-TW,zh;q=0.9,en-US;q=0.8,en;q=0.7",
            "Referer": origin_url,
        }

        try:
            try:
                response = requests.get(filepath, headers=headers, timeout=25)
            except (requests.exceptions.SSLError, requests.exceptions.ConnectionError):
                import urllib3
                urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
                response = requests.get(filepath, headers=headers, verify=False, timeout=25)
            response.raise_for_status()

            # 解析副檔名（去除 query string）
            url_path = unquote(urlparse(filepath).path)
            ext = os.path.splitext(url_path)[1].lower()
            content_type = response.headers.get("Content-Type", "").lower()

            # 從 Content-Disposition 取得真實檔名與副檔名
            cd = response.headers.get("Content-Disposition", "")
            if cd:
                m_star = re.search(r"filename\*\s*=\s*([^;']+''|UTF-8''|utf-8'')([^;\n]+)", cd, re.IGNORECASE)
                if m_star:
                    raw_fn = unquote(m_star.group(2).strip().strip('"\''))
                    ext_cd = os.path.splitext(raw_fn)[1].lower()
                    if ext_cd:
                        ext = ext_cd
                if not ext:
                    m_norm = re.search(r'filename\s*=\s*(["\'])(.*?)\1|filename\s*=\s*([^;\s\n]+)', cd, re.IGNORECASE)
                    if m_norm:
                        raw_fn = m_norm.group(2) or m_norm.group(3) or ""
                        ext_cd = os.path.splitext(raw_fn.strip())[1].lower()
                        if ext_cd:
                            ext = ext_cd

            # 若為網頁 (HTML)，直接提取乾淨內文與表格，避免混入原始 HTML 標籤
            # （注意：如果是 .xls 檔案回傳 text/html，不要直接當網頁，交給後續 Excel 處理器處理偽裝表格）
            if ext not in [".xls", ".xlsx", ".csv", ".pdf", ".docx", ".doc", ".json"]:
                if "text/html" in content_type or (not ext and "application/pdf" not in content_type) or ext in [".html", ".htm", ".php", ".asp", ".aspx", ".jsp"]:
                    try:
                        from tools.web_scraper import extract_clean_text_and_tables
                        response.encoding = response.apparent_encoding or "utf-8"
                        clean_text, _ = extract_clean_text_and_tables(response.text, base_url=filepath)
                        if clean_text and clean_text.strip():
                            return clean_text.strip()
                    except Exception:
                        pass

            # 從 Content-Type 推斷副檔名
            if not ext or ext in [".php", ".asp", ".aspx", ".jsp", ".cgi", ".bin"]:
                type_map = {
                    "application/pdf": ".pdf",
                    "application/msword": ".docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
                    "application/vnd.ms-excel": ".xls",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
                    "application/msexcel": ".xls",
                    "application/x-msexcel": ".xls",
                    "application/x-ms-excel": ".xls",
                    "application/x-excel": ".xls",
                    "application/xls": ".xls",
                    "text/plain": ".txt",
                    "text/csv": ".csv",
                    "application/json": ".json",
                }
                pure_ct = content_type.split(";")[0].strip()
                ext = type_map.get(pure_ct, ext or ".txt")

            cache_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "cache", "downloads")
            os.makedirs(cache_dir, exist_ok=True)
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=ext, dir=cache_dir)
            temp_file.write(response.content)
            temp_file.close()
            temp_file_path = temp_file.name
            filepath = temp_file.name
        except Exception as e:
            return f"Error: 無法下載線上檔案 - {e}"
    else:
        # 【防禦一】無論檔案是否存在，本機路徑都必須先通過沙盒白名單檢查
        # 這是防止 AI Tool Calling 傳入 ../../etc/passwd 等路徑的核心防線
        if not is_safe_local_path(filepath):
            return "錯誤：無權限讀取指定路徑之檔案。"
        if not os.path.exists(filepath):
            return "錯誤：找不到指定的文獻檔案。"

    ext = os.path.splitext(filepath)[1].lower()
    text = ""

    try:
        if ext == ".txt":
            text = _read_text_with_encodings(filepath)

        elif ext == ".pdf":
            reader = PdfReader(filepath)
            total_pages = len(reader.pages)
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text += f"\n[第 {i+1} 頁]\n" + page_text.strip() + "\n"
            if not text.strip():
                return f"注意：此 PDF 文件共有 {total_pages} 頁，但疑似為純影像掃描檔，未包含純文字層可供提取。"

        elif ext == ".docx":
            doc = Document(filepath)
            for para in doc.paragraphs:
                if para.text.strip():
                    text += para.text.strip() + "\n"

        elif ext == ".doc":
            # 嘗試 python-docx（若是命名為 .doc 的 docx 格式）
            try:
                doc = Document(filepath)
                for para in doc.paragraphs:
                    if para.text.strip():
                        text += para.text.strip() + "\n"
            except Exception:
                # 嘗試當作純文字/HTML
                try:
                    txt = _read_text_with_encodings(filepath)
                    if "<html" in txt.lower():
                        from bs4 import BeautifulSoup
                        soup = BeautifulSoup(txt, "html.parser")
                        text = soup.get_text(separator="\n", strip=True)
                    else:
                        # 二進位過濾文字
                        with open(filepath, "rb") as f:
                            raw_b = f.read()
                        ascii_strings = re.findall(rb"[\x20-\x7e\xa1-\xfe]{4,}", raw_b)
                        if ascii_strings:
                            text = "\n".join(s.decode("latin1", errors="ignore") for s in ascii_strings[:500])
                except Exception:
                    return "Error: 暫不支援舊版二進位 .doc 格式，建議轉為 .docx 後使用。"

        elif ext in [".xlsx", ".xls"]:
            text = _read_excel_file(filepath, ext)
            if text.startswith("Error:"):
                return text

        elif ext == ".csv":
            text = _read_csv_file(filepath)

        elif ext == ".json":
            raw_json = _read_text_with_encodings(filepath)
            data = json.loads(raw_json)
            text = json.dumps(data, ensure_ascii=False, indent=2)

        else:
            return f"Error: 不支援的檔案格式 {ext}"

        return text.strip()
    except PermissionError as e:
        if temp_file_path and os.path.exists(temp_file_path):
            try:
                os.remove(temp_file_path)
            except Exception:
                pass
        return mask_error(e, context="read_local_file/PermissionError")
    except Exception as e:
        # 【防禦三】錯誤遮蔽：僅將詳細資訊記錄到本機日誌，前端只顯示模糊訊息
        if temp_file_path and os.path.exists(temp_file_path):
            try:
                os.remove(temp_file_path)
            except Exception:
                pass
        return mask_error(e, context="read_local_file")
